from __future__ import annotations

from pathlib import Path

from docx import Document

from chunker import ChunkConfig, chunk_mevzuat
from file_reader import ReadConfig, read_docx


def create_test_docx(path: Path) -> None:
    document = Document()

    document.add_paragraph("MADDE 1 Ders yükü")
    document.add_paragraph(
        "Öğrencinin alabileceği dersler aşağıdaki tabloda gösterilir."
    )

    table = document.add_table(rows=3, cols=2)

    table.cell(0, 0).text = "Ders"
    table.cell(0, 1).text = "AKTS"

    table.cell(1, 0).text = "Yapay Zeka"
    table.cell(1, 1).text = "5 AKTS"

    table.cell(2, 0).text = "Veri Tabanı"
    table.cell(2, 1).text = "6 AKTS"

    document.add_paragraph("MADDE 2 Devam koşulu")
    document.add_paragraph(
        "Teorik derslere en az yüzde 70 devam zorunludur."
    )

    document.save(path)


def test_real_docx_reader_preserves_paragraph_and_table_order(
    #pytest’in hazır sunduğu bir fixture’dır.
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "ornek_yonetmelik.docx"
    create_test_docx(docx_path)

    extracted_text = read_docx(
        docx_path,
        cfg=ReadConfig(
            include_docx_table_markers=True,
        ),
    )

    assert "[DOCX_TABLE]" in extracted_text
    assert "[/DOCX_TABLE]" in extracted_text

    assert "Ders | AKTS" in extracted_text
    assert "Yapay Zeka | 5 AKTS" in extracted_text
    assert "Veri Tabanı | 6 AKTS" in extracted_text

    article_one_position = extracted_text.index(
        "MADDE 1 Ders yükü"
    )
    table_position = extracted_text.index(
        "[DOCX_TABLE]"
    )
    article_two_position = extracted_text.index(
        "MADDE 2 Devam koşulu"
    )

    assert article_one_position < table_position
    assert table_position < article_two_position


def test_real_docx_reader_to_chunker_integration(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "ornek_yonetmelik.docx"
    create_test_docx(docx_path)

    extracted_text = read_docx(
        docx_path,
        cfg=ReadConfig(
            include_docx_table_markers=True,
        ),
    )

    document = {
        "source": docx_path.name,
        "path": str(docx_path),
        "ext": ".docx",
        "size_bytes": docx_path.stat().st_size,
        "mtime": int(docx_path.stat().st_mtime),
        "text": extracted_text,
    }

    chunks = chunk_mevzuat(
        [document],
        cfg=ChunkConfig(
            chunk_size=5000,
            chunk_overlap=0,
            min_chunk_size=1,
        ),
    )

    assert len(chunks) == 2

    chunks_by_article = {
        chunk["madde_no"]: chunk
        for chunk in chunks
    }

    article_one = chunks_by_article[1]
    article_two = chunks_by_article[2]

    assert "[DOCX_TABLE]" not in article_one["text"]
    assert "[/DOCX_TABLE]" not in article_one["text"]

    assert "Yapay Zeka | 5 AKTS" in article_one["text"]
    assert "Veri Tabanı | 6 AKTS" in article_one["text"]

    assert article_one["contains_akts"] is True
    assert article_one["akts_values"] == [5, 6]

    assert "Teorik derslere en az yüzde 70" in article_two["text"]

    assert "Yapay Zeka | 5 AKTS" not in article_two["text"]
    assert "Veri Tabanı | 6 AKTS" not in article_two["text"]
